# -*- coding: utf-8 -*-
"""
爬虫控制面板后端
提供配置读写、爬虫启停、状态查询、实时日志推送
"""
import asyncio
import json
import os
import signal
import subprocess
import time
from typing import Dict, List, Optional

import logging

from fastapi import APIRouter, Request, WebSocket, WebSocketDisconnect
from fastapi.responses import JSONResponse

router = APIRouter(prefix="/control", tags=["control"])

# ── 项目根目录 & 配置文件路径 ──
PROJECT_ROOT = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
CONFIG_PATH = os.path.join(PROJECT_ROOT, "config", "anti_crawl_config.json")
DATA_DIR = os.path.join(PROJECT_ROOT, "data")
CRAWLER_LOG_PATH = os.path.join(PROJECT_ROOT, "logs", "crawler.log")

# ── 全局进程状态 ──
_crawler_process: Optional[asyncio.subprocess.Process] = None
_crawler_start_time: Optional[float] = None
_crawler_exit_code: Optional[int] = None
_log_lines: list = []            # 最近的日志行（环形缓冲区）
_LOG_MAX_LINES = 5000
_log_subscribers: list = []      # WebSocket 订阅者
_file_tail_task: Optional[asyncio.Task] = None  # 外部爬虫时 tail 日志文件


def _read_log_tail(n: int = 500) -> list:
    """从 crawler.log 读取最后 n 行（外部爬虫时用）"""
    if not os.path.exists(CRAWLER_LOG_PATH):
        return []
    try:
        with open(CRAWLER_LOG_PATH, "r", encoding="utf-8", errors="replace") as f:
            lines = f.readlines()
        return [line.rstrip("\n\r") for line in lines[-n:] if line.strip()]
    except Exception:
        return []


async def _file_tail_loop():
    """外部爬虫运行时，tail crawler.log 并推送给订阅者"""
    global _log_lines, _file_tail_task
    if not os.path.exists(CRAWLER_LOG_PATH):
        await asyncio.sleep(2)
        return
    last_ext_check = 0.0
    try:
        with open(CRAWLER_LOG_PATH, "r", encoding="utf-8", errors="replace") as f:
            f.seek(0, 2)  # 定位到文件末尾
            while _log_subscribers:
                if time.time() - last_ext_check > 5:
                    last_ext_check = time.time()
                    if not _detect_external_crawler():
                        break
                line = f.readline()
                if line:
                    text = line.rstrip("\n\r")
                    if text:
                        _log_lines.append(text)
                        if len(_log_lines) > _LOG_MAX_LINES:
                            _log_lines[:] = _log_lines[-_LOG_MAX_LINES:]
                        dead = []
                        for ws in _log_subscribers:
                            try:
                                await ws.send_text(json.dumps({"type": "log", "line": text}))
                            except Exception:
                                dead.append(ws)
                        for ws in dead:
                            if ws in _log_subscribers:
                                _log_subscribers.remove(ws)
                else:
                    await asyncio.sleep(0.3)
    except (asyncio.CancelledError, Exception):
        pass
    finally:
        _file_tail_task = None


def _ensure_file_tail_started():
    """有订阅者且为外部爬虫时，启动文件 tail 任务"""
    global _file_tail_task
    if (
        _log_subscribers
        and not _crawler_process
        and _detect_external_crawler()
        and (_file_tail_task is None or _file_tail_task.done())
    ):
        _file_tail_task = asyncio.create_task(_file_tail_loop())


# ================================================================
#  配置读写
# ================================================================

def _read_config() -> dict:
    """读取 anti_crawl_config.json"""
    with open(CONFIG_PATH, "r", encoding="utf-8") as f:
        return json.load(f)


def _write_config(data: dict):
    """写入 anti_crawl_config.json（保留格式）"""
    with open(CONFIG_PATH, "w", encoding="utf-8") as f:
        json.dump(data, f, ensure_ascii=False, indent=2)


@router.get("/config")
async def get_config():
    """获取所有配置"""
    try:
        cfg = _read_config()
        return {"success": True, "config": cfg}
    except Exception as e:
        return JSONResponse(status_code=500, content={"success": False, "error": str(e)})


@router.put("/config")
async def update_config(body: dict):
    """更新配置（整体覆写）"""
    try:
        _write_config(body)
        return {"success": True, "message": "配置已保存"}
    except Exception as e:
        return JSONResponse(status_code=500, content={"success": False, "error": str(e)})


@router.patch("/config")
async def patch_config(body: dict):
    """部分更新配置（合并到现有配置上层 key）"""
    try:
        cfg = _read_config()
        for key, value in body.items():
            if isinstance(value, dict) and isinstance(cfg.get(key), dict):
                cfg[key].update(value)
            else:
                cfg[key] = value
        _write_config(cfg)
        return {"success": True, "message": "配置已更新"}
    except Exception as e:
        return JSONResponse(status_code=500, content={"success": False, "error": str(e)})


# ================================================================
#  进程管理
# ================================================================

async def _stream_output(stream, label: str):
    """异步读取子进程输出并分发到日志"""
    global _log_lines
    try:
        while True:
            line = await stream.readline()
            if not line:
                break
            text = line.decode("utf-8", errors="replace").rstrip("\n")
            log_entry = f"{text}"
            _log_lines.append(log_entry)
            # 环形缓冲区
            if len(_log_lines) > _LOG_MAX_LINES:
                _log_lines = _log_lines[-_LOG_MAX_LINES:]
            # 推送给所有 WebSocket 订阅者
            dead = []
            for ws in _log_subscribers:
                try:
                    await ws.send_text(json.dumps({"type": "log", "line": log_entry}))
                except Exception:
                    dead.append(ws)
            for ws in dead:
                _log_subscribers.remove(ws)
    except Exception:
        pass


async def _wait_process():
    """等待进程退出并更新状态"""
    global _crawler_process, _crawler_exit_code
    if _crawler_process:
        code = await _crawler_process.wait()
        _crawler_exit_code = code
        # 通知订阅者进程已结束
        msg = json.dumps({"type": "status", "status": "stopped", "exit_code": code})
        dead = []
        for ws in _log_subscribers:
            try:
                await ws.send_text(msg)
            except Exception:
                dead.append(ws)
        for ws in dead:
            _log_subscribers.remove(ws)


def _build_batch_crawler_cmd() -> list:
    """根据 anti_crawl_config.json 构建 batch_crawler 完整命令"""
    cmd = [
        "conda", "run", "--no-capture-output", "-n", "uvenv",
        "python", "-u", "-m", "tools.batch_crawler",
    ]
    try:
        cfg = _read_config()
        batch = cfg.get("batch_crawl", {})
        feishu = cfg.get("feishu", {})

        excel_path = batch.get("excel_path", "")
        if excel_path:
            cmd.extend(["--excel", excel_path])

        max_notes = batch.get("max_notes_per_creator", 3000)
        cmd.extend(["--max-notes", str(max_notes)])

        if batch.get("enable_comments", False):
            cmd.append("--enable-comments")

        min_interaction = batch.get("min_interaction", 50)
        cmd.extend(["--min-interaction", str(min_interaction)])

        export_fmt = batch.get("export_format", "excel")
        cmd.extend(["--export-format", export_fmt])

        export_dir = batch.get("export_dir", "data/export")
        cmd.extend(["--export-dir", export_dir])

        if not feishu.get("enabled", False):
            cmd.append("--skip-feishu")
        else:
            for key, flag in [("app_id", "--feishu-app-id"), ("app_secret", "--feishu-app-secret"), ("folder_token", "--feishu-folder")]:
                val = feishu.get(key, "")
                if val:
                    cmd.extend([flag, val])

        crawl_mode = batch.get("crawl_mode", "full")
        if crawl_mode and crawl_mode != "full":
            cmd.extend(["--mode", crawl_mode])

        if not batch.get("resume", True):
            cmd.append("--no-resume")
        if batch.get("force_recrawl", False):
            cmd.append("--force-recrawl")

    except Exception:
        pass

    return cmd


def _get_task_info() -> dict:
    """读取当前任务摘要"""
    info = {"task_id": "", "total_creators": 0, "completed": 0, "failed": 0, "remaining": 0}
    try:
        cfg = _read_config()
        info["task_id"] = cfg.get("batch_crawl", {}).get("task_id", "")
    except Exception:
        pass
    if not info["task_id"]:
        from datetime import datetime as _dt
        info["task_id"] = _dt.now().strftime("task_%Y%m%d")

    try:
        excel_path = cfg.get("batch_crawl", {}).get("excel_path", "redbookaccontidandresult.xlsx")
        full_path = os.path.join(PROJECT_ROOT, excel_path) if not os.path.isabs(excel_path) else excel_path
        if os.path.exists(full_path):
            from tools.excel_reader import ExcelCreatorReader
            with ExcelCreatorReader(full_path) as reader:
                info["total_creators"] = len(reader.get_creators())
    except Exception:
        pass

    try:
        prog_file = os.path.join(PROJECT_ROOT, "data", f"batch_progress_{info['task_id']}.json")
        if os.path.exists(prog_file):
            with open(prog_file, "r", encoding="utf-8") as f:
                prog = json.load(f)
            info["completed"] = len(prog.get("completed", []))
            info["failed"] = len(prog.get("failed", {}))
            info["remaining"] = max(0, info["total_creators"] - info["completed"])
    except Exception:
        info["remaining"] = info["total_creators"]

    return info


@router.post("/start")
async def start_crawler():
    """启动爬虫进程"""
    global _crawler_process, _crawler_start_time, _crawler_exit_code, _log_lines

    if _crawler_process and _crawler_process.returncode is None:
        return JSONResponse(
            status_code=409,
            content={"success": False, "message": "爬虫正在运行中（本面板启动），请先停止再启动"}
        )

    ext = _detect_external_crawler()
    if ext:
        return JSONResponse(
            status_code=409,
            content={"success": False, "message": f"爬虫已在运行中（pid={ext['pid']}），请先停止再启动"}
        )

    _log_lines = []
    _crawler_exit_code = None

    cmd = _build_batch_crawler_cmd()
    task_info = _get_task_info()

    try:
        _crawler_process = await asyncio.create_subprocess_exec(
            *cmd,
            stdout=asyncio.subprocess.PIPE,
            stderr=asyncio.subprocess.STDOUT,
            cwd=PROJECT_ROOT,
            env={**os.environ, "PYTHONUNBUFFERED": "1"},
            start_new_session=True,
        )
        _crawler_start_time = time.time()

        cmd_display = " ".join(cmd[5:])
        mode_labels = {"full": "全量爬取", "incremental": "增量更新", "date_range": "日期范围"}
        try:
            _cfg_mode = _read_config().get("batch_crawl", {}).get("crawl_mode", "full")
        except Exception:
            _cfg_mode = "full"
        _log_lines.append(f"[控制面板] 任务: {task_info['task_id']} | 模式: {mode_labels.get(_cfg_mode, _cfg_mode)} | 总计: {task_info['total_creators']}个作者 | 已完成: {task_info['completed']} | 待爬取: {task_info['remaining']}")
        _log_lines.append(f"[控制面板] 命令: python -u -m {cmd_display}")

        asyncio.create_task(_stream_output(_crawler_process.stdout, "stdout"))
        asyncio.create_task(_wait_process())

        return {
            "success": True,
            "message": f"任务 {task_info['task_id']} 已启动 | 待爬取 {task_info['remaining']}/{task_info['total_creators']} 个作者",
            "pid": _crawler_process.pid,
        }
    except Exception as e:
        return JSONResponse(
            status_code=500,
            content={"success": False, "error": f"启动失败: {str(e)}"}
        )


def _log_stop_audit(source: str, request: Request, extra: str = ""):
    client = getattr(request, "client", None)
    ip = client.host if client else "unknown"
    forwarded = request.headers.get("x-forwarded-for", "").split(",")[0].strip() or ip
    ua = request.headers.get("user-agent", "")[:80]
    logging.getLogger("MediaCrawler").info(
        f"[StopAudit] {source} | 来源={forwarded} | UA={ua} | {extra}"
    )


@router.post("/stop")
async def stop_crawler(request: Request):
    """优雅停止爬虫进程（SIGTERM → 等待数据保存 → 退出）"""
    global _crawler_process

    _log_stop_audit("控制面板停止请求", request)

    # 面板启动的进程（通过进程组确保 conda 包装和实际 python 进程都收到信号）
    if _crawler_process and _crawler_process.returncode is None:
        try:
            pid = _crawler_process.pid
            _log_lines.append("[控制面板] 正在优雅停止，保存当前进度...")
            try:
                pgid = os.getpgid(pid)
                os.killpg(pgid, signal.SIGTERM)
            except OSError:
                _crawler_process.terminate()
            graceful = False
            try:
                await asyncio.wait_for(_crawler_process.wait(), timeout=30)
                graceful = True
            except asyncio.TimeoutError:
                try:
                    pgid = os.getpgid(pid)
                    os.killpg(pgid, signal.SIGKILL)
                except OSError:
                    _crawler_process.kill()
                await _crawler_process.wait()
            msg = f"爬虫已优雅停止 (pid={pid})" if graceful else f"爬虫已强制停止 (pid={pid})"
            _log_lines.append(f"[控制面板] {msg}")
            return {"success": True, "message": msg, "graceful": graceful}
        except Exception as e:
            return JSONResponse(
                status_code=500,
                content={"success": False, "error": str(e)}
            )

    # 外部启动的进程
    ext = _detect_external_crawler()
    if ext:
        try:
            pids = ext["pids"]
            _log_lines.append(f"[控制面板] 正在停止外部爬虫进程 (pids={pids})...")
            for pid in pids:
                try:
                    os.kill(pid, signal.SIGTERM)
                except OSError:
                    pass
            graceful = False
            for _ in range(30):
                await asyncio.sleep(1)
                alive = [p for p in pids if _pid_alive(p)]
                if not alive:
                    graceful = True
                    break
            if not graceful:
                for pid in pids:
                    try:
                        os.kill(pid, signal.SIGKILL)
                    except OSError:
                        pass
                await asyncio.sleep(1)
            msg = f"爬虫已优雅停止 (pids={pids})" if graceful else f"爬虫已强制停止 (pids={pids})"
            _log_lines.append(f"[控制面板] {msg}")
            return {"success": True, "message": msg, "graceful": graceful}
        except Exception as e:
            return JSONResponse(
                status_code=500,
                content={"success": False, "error": str(e)}
            )

    return {"success": False, "message": "爬虫未在运行"}


def _pid_alive(pid: int) -> bool:
    try:
        os.kill(pid, 0)
        return True
    except OSError:
        return False


def _detect_external_crawler() -> dict | None:
    """检测是否有外部启动的 batch_crawler 进程在运行（返回所有相关 PID）"""
    import subprocess
    try:
        result = subprocess.run(
            ["pgrep", "-f", "tools.batch_crawler"],
            capture_output=True, text=True, timeout=5,
        )
        if result.returncode != 0:
            return None
        pids = [int(p) for p in result.stdout.strip().split("\n") if p.strip()]
        if not pids:
            return None
        # 优先取实际 python 进程的运行时间
        main_pid = pids[-1]
        elapsed = 0
        try:
            stat_path = f"/proc/{main_pid}/stat"
            if os.path.exists(stat_path):
                with open(stat_path) as f:
                    fields = f.read().split()
                starttime_ticks = int(fields[21])
                with open("/proc/uptime") as f:
                    uptime_sec = float(f.read().split()[0])
                clk_tck = os.sysconf("SC_CLK_TCK")
                start_sec = starttime_ticks / clk_tck
                elapsed = int(uptime_sec - start_sec)
        except Exception:
            pass
        return {"pid": main_pid, "pids": pids, "running_seconds": elapsed}
    except Exception:
        return None


@router.get("/status")
async def get_status():
    """获取爬虫运行状态"""
    if _crawler_process and _crawler_process.returncode is None:
        elapsed = int(time.time() - _crawler_start_time) if _crawler_start_time else 0
        return {
            "status": "running",
            "message": f"爬虫运行中 ({elapsed}s)",
            "pid": _crawler_process.pid,
            "running_seconds": elapsed,
            "exit_code": None,
        }

    ext = _detect_external_crawler()
    if ext:
        return {
            "status": "running",
            "message": f"爬虫运行中（外部启动, {ext['running_seconds']}s）",
            "pid": ext["pid"],
            "running_seconds": ext["running_seconds"],
            "exit_code": None,
            "external": True,
        }

    if _crawler_process is None:
        return {
            "status": "idle",
            "message": "爬虫未在运行",
            "pid": None,
            "running_seconds": 0,
            "exit_code": None,
        }

    return {
        "status": "stopped",
        "message": "爬虫已停止",
        "pid": None,
        "running_seconds": 0,
        "exit_code": _crawler_exit_code,
    }


@router.get("/logs")
async def get_logs(last: int = 200):
    """获取最近的日志行（优先内存，外部爬虫时从文件读）"""
    n = min(last, _LOG_MAX_LINES)
    lines = _log_lines[-n:] if _log_lines else []
    if not lines and _detect_external_crawler():
        lines = _read_log_tail(n)
    return {"lines": lines}


# ================================================================
#  WebSocket 实时日志
# ================================================================

@router.websocket("/ws/log")
async def ws_log(ws: WebSocket):
    """WebSocket 实时推送爬虫日志"""
    await ws.accept()
    _log_subscribers.append(ws)

    try:
        # 先发送历史日志：有则用 _log_lines，否则从 crawler.log 读取（外部爬虫场景）
        history = _log_lines[-500:] if _log_lines else []
        if not history and (_crawler_process or _detect_external_crawler()):
            history = _read_log_tail(500)
        if history:
            await ws.send_text(json.dumps({
                "type": "history",
                "lines": history,
            }))

        # 外部爬虫时启动文件 tail 任务
        _ensure_file_tail_started()

        # 发送当前状态（与 get_status 逻辑一致，含外部启动检测）
        if _crawler_process and _crawler_process.returncode is None:
            elapsed = int(time.time() - _crawler_start_time) if _crawler_start_time else 0
            await ws.send_text(json.dumps({
                "type": "status",
                "status": "running",
                "pid": _crawler_process.pid,
                "running_seconds": elapsed,
            }))
        else:
            ext = _detect_external_crawler()
            if ext:
                await ws.send_text(json.dumps({
                    "type": "status",
                    "status": "running",
                    "pid": ext["pid"],
                    "running_seconds": ext["running_seconds"],
                }))
            else:
                await ws.send_text(json.dumps({
                    "type": "status",
                    "status": "stopped",
                    "exit_code": _crawler_exit_code,
                }))

        # 保持连接存活
        while True:
            try:
                msg = await asyncio.wait_for(ws.receive_text(), timeout=30)
                # 客户端可发送 ping 保活
                if msg == "ping":
                    await ws.send_text(json.dumps({"type": "pong"}))
            except asyncio.TimeoutError:
                # 发送心跳
                try:
                    await ws.send_text(json.dumps({"type": "ping"}))
                except Exception:
                    break
    except WebSocketDisconnect:
        pass
    except Exception:
        pass
    finally:
        if ws in _log_subscribers:
            _log_subscribers.remove(ws)


# ================================================================
#  作者管理（Excel 读写）
# ================================================================

def _resolve_excel_path() -> str:
    try:
        cfg = _read_config()
        rel = cfg.get("batch_crawl", {}).get("excel_path", "redbookaccontidandresult.xlsx")
    except Exception:
        rel = "redbookaccontidandresult.xlsx"
    if os.path.isabs(rel):
        return rel
    return os.path.join(PROJECT_ROOT, rel)


def _read_creators() -> list:
    path = _resolve_excel_path()
    if not os.path.exists(path):
        return []
    from tools.excel_reader import ExcelCreatorReader
    with ExcelCreatorReader(path) as reader:
        return reader.get_creators()


def _write_creators(creators: list):
    """将作者列表写回 Excel（覆盖 '小红书账号' sheet，保留其他 sheet）"""
    import openpyxl
    path = _resolve_excel_path()
    if os.path.exists(path):
        wb = openpyxl.load_workbook(path)
    else:
        wb = openpyxl.Workbook()

    sheet_name = "小红书账号"
    if sheet_name in wb.sheetnames:
        del wb[sheet_name]
    ws = wb.create_sheet(sheet_name, 0)

    ws.append(["名称", "ID", "主页链接"])
    for c in creators:
        ws.append([c.get("name", ""), c.get("id", ""), c.get("url", "")])

    wb.save(path)
    wb.close()


@router.get("/creators")
async def list_creators():
    try:
        creators = _read_creators()
        return {"success": True, "creators": creators, "total": len(creators)}
    except Exception as e:
        return JSONResponse(status_code=500, content={"success": False, "error": str(e)})


@router.post("/creators")
async def add_creator(body: dict):
    try:
        name = body.get("name", "").strip()
        cid = body.get("id", "").strip()
        url = body.get("url", "").strip()
        if not url or "xiaohongshu.com" not in url:
            return JSONResponse(status_code=400, content={"success": False, "error": "链接必须包含 xiaohongshu.com"})

        creators = _read_creators()
        existing_urls = {c["url"].split("?")[0] for c in creators}
        clean_url = url.split("?")[0]
        if clean_url in existing_urls:
            return JSONResponse(status_code=409, content={"success": False, "error": "该作者已存在"})

        creators.append({"name": name, "id": cid, "url": url})
        _write_creators(creators)
        return {"success": True, "message": f"已添加: {name or url}", "total": len(creators)}
    except Exception as e:
        return JSONResponse(status_code=500, content={"success": False, "error": str(e)})


@router.put("/creators/{index}")
async def update_creator(index: int, body: dict):
    try:
        creators = _read_creators()
        if index < 0 or index >= len(creators):
            return JSONResponse(status_code=404, content={"success": False, "error": "索引无效"})

        url = body.get("url", "").strip()
        if url and "xiaohongshu.com" not in url:
            return JSONResponse(status_code=400, content={"success": False, "error": "链接必须包含 xiaohongshu.com"})

        if "name" in body:
            creators[index]["name"] = body["name"].strip()
        if "id" in body:
            creators[index]["id"] = body["id"].strip()
        if "url" in body and url:
            creators[index]["url"] = url

        _write_creators(creators)
        return {"success": True, "message": "已更新", "creator": creators[index]}
    except Exception as e:
        return JSONResponse(status_code=500, content={"success": False, "error": str(e)})


@router.delete("/creators/{index}")
async def delete_creator(index: int):
    try:
        creators = _read_creators()
        if index < 0 or index >= len(creators):
            return JSONResponse(status_code=404, content={"success": False, "error": "索引无效"})

        removed = creators.pop(index)
        _write_creators(creators)
        return {"success": True, "message": f"已删除: {removed.get('name', '')}", "total": len(creators)}
    except Exception as e:
        return JSONResponse(status_code=500, content={"success": False, "error": str(e)})


@router.put("/creators")
async def replace_all_creators(body: dict):
    """批量替换所有作者"""
    try:
        creators = body.get("creators", [])
        valid = []
        for c in creators:
            url = c.get("url", "").strip()
            if url and "xiaohongshu.com" in url:
                valid.append({"name": c.get("name", "").strip(), "id": c.get("id", "").strip(), "url": url})
        _write_creators(valid)
        return {"success": True, "message": f"已保存 {len(valid)} 个作者", "total": len(valid)}
    except Exception as e:
        return JSONResponse(status_code=500, content={"success": False, "error": str(e)})


# ================================================================
#  Session 池管理
# ================================================================

def _get_session_pool():
    from tools.session_pool import SessionPool
    return SessionPool()


@router.get("/sessions")
async def list_sessions():
    try:
        pool = _get_session_pool()
        return {"success": True, "sessions": pool.get_all(), "stats": pool.stats()}
    except Exception as e:
        return JSONResponse(status_code=500, content={"success": False, "error": str(e)})


@router.post("/sessions")
async def add_session(body: dict):
    try:
        ws = body.get("web_session", "").strip()
        label = body.get("label", "").strip()
        if not ws or len(ws) < 10:
            return JSONResponse(status_code=400, content={"success": False, "error": "web_session 无效"})
        pool = _get_session_pool()
        entry = pool.add(ws, label)
        return {"success": True, "message": f"已添加: {entry.id}", "session": entry.to_dict()}
    except ValueError as ve:
        return JSONResponse(status_code=409, content={"success": False, "error": str(ve)})
    except Exception as e:
        return JSONResponse(status_code=500, content={"success": False, "error": str(e)})


@router.delete("/sessions/{session_id}")
async def delete_session(session_id: str):
    try:
        pool = _get_session_pool()
        ok = pool.remove(session_id)
        if not ok:
            return JSONResponse(status_code=404, content={"success": False, "error": "未找到"})
        return {"success": True, "message": "已删除"}
    except Exception as e:
        return JSONResponse(status_code=500, content={"success": False, "error": str(e)})


@router.post("/sessions/{session_id}/verify")
async def verify_single_session(session_id: str):
    try:
        pool = _get_session_pool()
        entry = pool.get_by_id(session_id)
        if not entry:
            return JSONResponse(status_code=404, content={"success": False, "error": "未找到"})
        valid = await _verify_web_session(entry.web_session)
        if valid:
            pool.mark_active(session_id)
            return {"success": True, "status": "active", "message": "验证通过"}
        else:
            pool.mark_expired(session_id)
            return {"success": True, "status": "expired", "message": "已过期"}
    except Exception as e:
        return JSONResponse(status_code=500, content={"success": False, "error": str(e)})


@router.post("/sessions/verify-all")
async def verify_all_sessions():
    try:
        pool = _get_session_pool()
        results = {"verified": 0, "active": 0, "expired": 0}
        for s in pool.get_all():
            if s["status"] == "expired":
                continue
            valid = await _verify_web_session(s["web_session"])
            if valid:
                pool.mark_active(s["id"])
                results["active"] += 1
            else:
                pool.mark_expired(s["id"])
                results["expired"] += 1
            results["verified"] += 1
        return {"success": True, "results": results, "stats": pool.stats()}
    except Exception as e:
        return JSONResponse(status_code=500, content={"success": False, "error": str(e)})


async def _verify_web_session(web_session: str) -> bool:
    """轻量验证：用 httpx 发请求检查 session 是否有效"""
    import httpx
    try:
        headers = {
            "User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/120.0.0.0 Safari/537.36",
            "Cookie": f"web_session={web_session}",
            "Origin": "https://www.xiaohongshu.com",
            "Referer": "https://www.xiaohongshu.com/",
        }
        async with httpx.AsyncClient(timeout=15, follow_redirects=True) as client:
            resp = await client.get(
                "https://edith.xiaohongshu.com/api/sns/web/v1/user/selfinfo",
                headers=headers,
            )
            if resp.status_code == 200:
                data = resp.json()
                return data.get("success", False) or data.get("code", -1) == 0
    except Exception:
        pass
    return False


# ── 飞书重新写入（批量作者数据 → 飞书多维表格） ──

_resync_feishu_proc: Optional[subprocess.Popen] = None
_resync_feishu_logs: List[str] = []
_resync_feishu_status: str = ""  # "", "running", "success", "failed"


@router.post("/resync-feishu")
async def resync_feishu_creator():
    """一键将本地已爬取的作者数据重新写入飞书多维表格"""
    global _resync_feishu_proc, _resync_feishu_logs, _resync_feishu_status

    if _resync_feishu_proc and _resync_feishu_proc.poll() is None:
        return JSONResponse(
            status_code=409,
            content={"success": False, "message": "飞书写入正在运行中，请等待完成"}
        )

    cfg = _read_config()
    feishu = cfg.get("feishu", {})
    if not feishu.get("enabled"):
        return JSONResponse(
            status_code=400,
            content={"success": False, "message": "飞书推送未启用，请先开启"}
        )

    _resync_feishu_logs = []
    _resync_feishu_status = "running"

    cmd = [
        "conda", "run", "--no-capture-output", "-n", "uvenv",
        "--cwd", PROJECT_ROOT,
        "python", "-u", "-m", "tools.resync_feishu_creator",
    ]

    try:
        _resync_feishu_proc = subprocess.Popen(
            cmd,
            stdout=subprocess.PIPE,
            stderr=subprocess.STDOUT,
            text=True,
            encoding="utf-8",
            errors="replace",
            cwd=PROJECT_ROOT,
            env={**os.environ, "PYTHONUNBUFFERED": "1"},
        )
    except Exception as e:
        _resync_feishu_status = "failed"
        return JSONResponse(
            status_code=500,
            content={"success": False, "message": f"启动失败: {e}"}
        )

    async def _wait_resync():
        global _resync_feishu_status, _resync_feishu_proc
        loop = asyncio.get_event_loop()
        try:
            while _resync_feishu_proc and _resync_feishu_proc.poll() is None:
                line = await loop.run_in_executor(
                    None, _resync_feishu_proc.stdout.readline
                )
                if line:
                    stripped = line.strip()
                    if stripped:
                        _resync_feishu_logs.append(stripped)
                        if len(_resync_feishu_logs) > 2000:
                            _resync_feishu_logs[:] = _resync_feishu_logs[-1500:]
            remaining = await loop.run_in_executor(
                None, _resync_feishu_proc.stdout.read
            )
            if remaining:
                for ln in remaining.strip().split("\n"):
                    if ln.strip():
                        _resync_feishu_logs.append(ln.strip())
        except Exception:
            pass

        exit_code = _resync_feishu_proc.returncode if _resync_feishu_proc else -1
        full = "\n".join(_resync_feishu_logs)
        if "飞书写入完成" in full or exit_code == 0:
            _resync_feishu_status = "success"
        else:
            _resync_feishu_status = "failed"
        _resync_feishu_logs.append(
            f"=== 写入结束: exit_code={exit_code}, status={_resync_feishu_status} ==="
        )

    asyncio.create_task(_wait_resync())
    return {"success": True, "message": "飞书写入已启动，请在页面查看进度"}


@router.get("/resync-feishu/status")
async def resync_feishu_status():
    """获取飞书重新写入的状态和日志"""
    running = bool(
        _resync_feishu_proc and _resync_feishu_proc.poll() is None
    )
    bitable_url = ""
    for ln in reversed(_resync_feishu_logs):
        if "多维表格链接:" in ln:
            idx = ln.find("http")
            if idx >= 0:
                bitable_url = ln[idx:].strip()
            break

    return {
        "success": True,
        "status": "running" if running else _resync_feishu_status,
        "logs": _resync_feishu_logs[-200:],
        "bitable_url": bitable_url,
    }


@router.post("/resync-feishu/stop")
async def resync_feishu_stop(request: Request):
    """停止飞书写入进程"""
    global _resync_feishu_proc, _resync_feishu_status
    if not _resync_feishu_proc or _resync_feishu_proc.poll() is not None:
        return JSONResponse(
            status_code=400,
            content={"success": False, "message": "没有正在运行的飞书写入进程"}
        )
    try:
        _resync_feishu_proc.send_signal(signal.SIGTERM)
        _resync_feishu_proc.wait(timeout=10)
    except subprocess.TimeoutExpired:
        _resync_feishu_proc.kill()
    except Exception:
        pass
    _resync_feishu_status = "failed"
    _resync_feishu_logs.append("=== 已手动停止 ===")
    return {"success": True, "message": "飞书写入已停止"}
