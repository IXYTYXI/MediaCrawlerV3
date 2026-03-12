# -*- coding: utf-8 -*-
"""
将本地已爬取的「批量爬取（作者模式）」数据重新写入飞书多维表格。

用于：爬取已完成但飞书写入失败（如 TextFieldConvFail）时，
不重新爬取，仅用修好的代码把 data/xhs/json/{task_id}/ 下的
creator_{user_id}_contents.json 和 _reuse.json 再推一次到飞书。

用法:
  cd /path/to/MediaCrawler-stable
  # 需使用与爬虫相同的环境（含 playwright、openpyxl 等）
  conda run -n uvenv python -m tools.resync_feishu_creator                    # 使用 config 中的 task_id
  conda run -n uvenv python -m tools.resync_feishu_creator --task-id task_20260210
  conda run -n uvenv python -m tools.resync_feishu_creator --task-id task_20260210 --excel 其他.xlsx
"""
import json
import os
import sys
from datetime import datetime
from typing import Dict, List, Any, Optional

# 项目根目录
root = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, root)
os.chdir(root)


def _safe_int(v) -> int:
    if v is None or v == "":
        return 0
    try:
        return int(v)
    except (ValueError, TypeError):
        return 0


def _get_interaction_count(note: Dict) -> int:
    """计算笔记的互动量 = 点赞 + 评论 + 收藏"""
    liked = _safe_int(note.get("liked_count", 0))
    comment = _safe_int(note.get("comment_count", 0))
    collected = _safe_int(note.get("collected_count", 0))
    return liked + comment + collected


def _check_date_range(note: Dict, date_start: str, date_end: str) -> bool:
    """检查作品发布时间是否在指定日期范围内"""
    if not date_start and not date_end:
        return True
    time_val = note.get("time", 0)
    if not time_val or not isinstance(time_val, (int, float)):
        return True
    try:
        ts = time_val / 1000 if time_val > 1e12 else time_val
        publish_date = datetime.fromtimestamp(ts)
        if date_start:
            start = datetime.strptime(date_start, "%Y-%m-%d")
            if publish_date < start:
                return False
        if date_end:
            end = datetime.strptime(date_end, "%Y-%m-%d").replace(
                hour=23, minute=59, second=59
            )
            if publish_date > end:
                return False
        return True
    except Exception:
        return True


def _note_completeness(note: Dict) -> int:
    """评估笔记数据完整度，用于去重时保留更完整的"""
    score = 0
    if note.get("title"):
        score += 10
    if note.get("desc"):
        score += 5
    if note.get("image_list") or note.get("video_url"):
        score += 20
    if note.get("liked_count") is not None or note.get("collected_count") is not None:
        score += 5
    return score


def _load_creator_notes(
    task_dir: str,
    user_id: str,
    crawl_mode: str = "full",
    date_start: str = "",
    date_end: str = "",
    min_interaction: int = 0,
) -> List[Dict]:
    """
    读取单个作者在 task_dir 下的所有数据（contents + reuse），
    合并去重后按日期/互动量过滤。
    """
    try:
        from tools.note_merge import merge_note_metadata
    except ImportError:
        merge_note_metadata = lambda a, b: a if _note_completeness(a) >= _note_completeness(b) else b

    notes_by_id: Dict[str, Dict] = {}
    for suffix in ["_contents.json", "_reuse.json"]:
        fp = os.path.join(task_dir, f"creator_{user_id}{suffix}")
        if not os.path.exists(fp):
            continue
        try:
            with open(fp, "r", encoding="utf-8") as f:
                arr = json.load(f)
            if not isinstance(arr, list):
                continue
            for item in arr:
                nid = item.get("note_id", "")
                if not nid:
                    continue
                if nid not in notes_by_id:
                    notes_by_id[nid] = item
                else:
                    cur = notes_by_id[nid]
                    base = item if _note_completeness(item) > _note_completeness(cur) else cur
                    update = cur if base is item else item
                    notes_by_id[nid] = merge_note_metadata(base, update)
        except Exception as e:
            print(f"  [ResyncFeishu] 读取失败 {os.path.basename(fp)}: {e}")

    result = []
    for item in notes_by_id.values():
        if item.get("_reused"):
            result.append(item)
            continue
        if crawl_mode == "date_range" and not _check_date_range(item, date_start, date_end):
            continue
        if min_interaction > 0 and _get_interaction_count(item) < min_interaction:
            continue
        result.append(item)
    return result


def _load_creator_names_from_excel(excel_path: str) -> Dict[str, str]:
    """从 Excel 读取作者映射 {user_id: creator_name}，避免依赖 playwright"""
    import openpyxl

    user_id_to_name: Dict[str, str] = {}
    wb = openpyxl.load_workbook(excel_path, read_only=True)
    sheet_name = "小红书账号" if "小红书账号" in wb.sheetnames else wb.sheetnames[0]
    ws = wb[sheet_name]
    headers = [str(c).strip() if c else "" for c in next(ws.iter_rows(min_row=1, max_row=1, values_only=True))]
    col_name = next((i for i, h in enumerate(headers) if "名称" in h or "name" in h.lower()), 0)
    col_url = next((i for i, h in enumerate(headers) if "链接" in h or "url" in h.lower() or "主页" in h), -1)
    if col_url < 0:
        return user_id_to_name
    for row in ws.iter_rows(min_row=2, values_only=True):
        row_l = list(row)
        url = str(row_l[col_url]).strip() if col_url < len(row_l) and row_l[col_url] else ""
        if not url or "xiaohongshu.com" not in url:
            continue
        name = str(row_l[col_name]).strip() if col_name < len(row_l) and row_l[col_name] else "未知"
        if "/user/profile/" in url:
            uid = url.split("/user/profile/")[1].split("?")[0].split("/")[0]
            user_id_to_name[uid] = name
    wb.close()
    return user_id_to_name


def _collect_creators_with_data(
    task_dir: str,
    excel_path: str,
    crawl_mode: str,
    date_start: str,
    date_end: str,
    min_interaction: int,
) -> List[tuple]:
    """
    收集有数据的作者列表 (creator_name, user_id, notes)。
    从 task_dir 扫描 creator_*_contents.json 得到 user_id，
    从 Excel 获取 creator_name，加载并过滤 notes。
    """
    # 1. 从 Excel 读取作者映射 {user_id: creator_name}
    user_id_to_name = _load_creator_names_from_excel(excel_path)

    # 2. 扫描 task_dir 获取有数据的 user_id
    if not os.path.isdir(task_dir):
        print(f"[ResyncFeishu] 任务目录不存在: {task_dir}")
        return []

    collected: List[tuple] = []
    for name in os.listdir(task_dir):
        if not name.endswith("_contents.json") or not name.startswith("creator_"):
            continue
        uid = name[len("creator_") : -len("_contents.json")]
        creator_name = user_id_to_name.get(uid, uid[:12] + "...")
        notes = _load_creator_notes(
            task_dir, uid, crawl_mode, date_start, date_end, min_interaction
        )
        if notes:
            collected.append((creator_name, uid, notes))
            print(f"  [ResyncFeishu] {creator_name}: {len(notes)} 条")

    return collected


def main():
    import argparse

    parser = argparse.ArgumentParser(
        description="批量爬取数据重新写入飞书（不重新爬取）"
    )
    parser.add_argument(
        "--task-id",
        default="",
        help="任务ID，默认从 config.batch_crawl.task_id 读取",
    )
    parser.add_argument(
        "--excel",
        default="",
        help="Excel 作者列表路径，默认从 config.batch_crawl.excel_path 读取",
    )
    parser.add_argument(
        "--folder",
        default="",
        help="飞书文件夹 token，默认从 config.feishu.folder_token 读取",
    )
    args = parser.parse_args()

    # 读取配置
    cfg_path = os.path.join("config", "anti_crawl_config.json")
    try:
        with open(cfg_path, "r", encoding="utf-8") as f:
            cfg = json.load(f)
    except Exception as e:
        print(f"[ResyncFeishu] 读取配置失败: {e}")
        sys.exit(1)

    batch_cfg = cfg.get("batch_crawl", {}) or {}
    feishu_cfg = cfg.get("feishu", {}) or {}

    task_id = args.task_id or batch_cfg.get("task_id", "")
    if not task_id:
        task_id = "task_20260210"
        print(f"[ResyncFeishu] 未配置 task_id，使用默认: {task_id}")

    excel_path = args.excel or batch_cfg.get("excel_path", "redbookaccontidandresult.xlsx")
    if not os.path.isfile(excel_path):
        print(f"[ResyncFeishu] Excel 文件不存在: {excel_path}")
        sys.exit(1)

    if not feishu_cfg.get("enabled"):
        print("[ResyncFeishu] 飞书未启用，请在 config 中开启")
        sys.exit(1)

    app_id = feishu_cfg.get("app_id", "")
    app_secret = feishu_cfg.get("app_secret", "")
    if not app_id or not app_secret:
        print("[ResyncFeishu] 未配置 feishu.app_id / app_secret")
        sys.exit(1)

    folder_token = args.folder or feishu_cfg.get("folder_token", "")
    reuse_token = ""
    if feishu_cfg.get("reuse_bitable"):
        lb = feishu_cfg.get("last_bitable", {})
        if isinstance(lb, dict):
            reuse_token = lb.get("app_token", "")

    crawl_mode = batch_cfg.get("crawl_mode", "full")
    date_start = batch_cfg.get("date_start", "")
    date_end = batch_cfg.get("date_end", "")
    min_interaction = int(batch_cfg.get("min_interaction", 50))

    task_dir = os.path.join("data", "xhs", "json", task_id)
    print(f"[ResyncFeishu] 任务目录: {task_dir}")
    print(f"[ResyncFeishu] Excel: {excel_path}")
    print(f"[ResyncFeishu] 互动量过滤: {'关闭' if min_interaction <= 0 else f'>={min_interaction}'}")
    if reuse_token:
        print(f"[ResyncFeishu] 复用已有多维表格: {reuse_token[:20]}...")
    print("-" * 50)

    creators_data = _collect_creators_with_data(
        task_dir, excel_path, crawl_mode, date_start, date_end, min_interaction
    )

    if not creators_data:
        print("[ResyncFeishu] 没有可写入的作者数据")
        sys.exit(0)

    total_notes = sum(len(notes) for _, _, notes in creators_data)
    print(f"\n[ResyncFeishu] 共 {len(creators_data)} 个作者, {total_notes} 条笔记，开始写入飞书...")

    owner_open_id = feishu_cfg.get("owner_open_id", "")

    try:
        from tools.pipeline_feishu_writer import PipelineFeishuWriter

        with PipelineFeishuWriter(
            feishu_app_id=app_id,
            feishu_app_secret=app_secret,
            folder_token=folder_token,
            reuse_app_token=reuse_token,
        ) as writer:
            if owner_open_id:
                writer._owner_open_id = owner_open_id
            for creator_name, _uid, notes in creators_data:
                writer.write_creator(creator_name, notes)
            writer.finalize()

        print("[ResyncFeishu] 飞书写入完成")
        if hasattr(writer, "bitable_url") and writer.bitable_url:
            print(f"[ResyncFeishu] 多维表格链接: {writer.bitable_url}")
    except Exception as e:
        print(f"[ResyncFeishu] 飞书写入失败: {e}")
        import traceback
        traceback.print_exc()
        sys.exit(1)


if __name__ == "__main__":
    main()
