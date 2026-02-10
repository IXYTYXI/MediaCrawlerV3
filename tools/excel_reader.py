# -*- coding: utf-8 -*-
"""
Excel 作者列表读取器
从本地 Excel 文件读取小红书作者列表和结果字段定义
"""
import os
from typing import List, Dict, Optional, Tuple

import openpyxl

from tools import utils


class ExcelCreatorReader:
    """从 Excel 文件读取作者列表"""

    def __init__(self, file_path: str):
        """
        Args:
            file_path: Excel 文件路径
        """
        if not os.path.exists(file_path):
            raise FileNotFoundError(f"Excel 文件不存在: {file_path}")
        self.file_path = file_path
        self._wb = openpyxl.load_workbook(file_path, read_only=True)

    def get_creators(self, sheet_name: str = "小红书账号") -> List[Dict[str, str]]:
        """
        读取作者列表
        
        Args:
            sheet_name: Sheet 名称，默认 "小红书账号"
            
        Returns:
            [{"name": "xxx", "id": "xxx", "url": "https://..."}]
        """
        if sheet_name not in self._wb.sheetnames:
            raise ValueError(f"Sheet '{sheet_name}' 不存在, 可用: {self._wb.sheetnames}")

        ws = self._wb[sheet_name]
        creators = []

        # 读取表头
        headers = []
        for cell in next(ws.iter_rows(min_row=1, max_row=1, values_only=True)):
            headers.append(str(cell).strip() if cell else "")

        # 映射列名
        col_map = {}
        for idx, h in enumerate(headers):
            if "名称" in h or "name" in h.lower():
                col_map["name"] = idx
            elif h.upper() == "ID" or ("账号" in h and "链接" not in h):
                col_map["id"] = idx
            elif "链接" in h or "url" in h.lower() or "主页" in h:
                col_map["url"] = idx

        if "url" not in col_map:
            raise ValueError(f"未找到包含'链接'或'url'的列, 表头: {headers}")

        # 读取数据行
        for row in ws.iter_rows(min_row=2, values_only=True):
            row_list = list(row)
            url = str(row_list[col_map["url"]]).strip() if col_map.get("url") is not None and row_list[col_map["url"]] else ""
            if not url or "xiaohongshu.com" not in url:
                continue

            creator = {
                "name": str(row_list[col_map.get("name", 0)]).strip() if col_map.get("name") is not None and row_list[col_map.get("name", 0)] else "",
                "id": str(row_list[col_map.get("id", 1)]).strip() if col_map.get("id") is not None and row_list[col_map.get("id", 1)] else "",
                "url": url,
            }
            creators.append(creator)

        utils.logger.info(f"[ExcelReader] 从 '{sheet_name}' 读取到 {len(creators)} 个作者")
        return creators

    def get_result_fields(self, sheet_name: str = "需要的结果字段") -> List[Dict[str, str]]:
        """
        读取结果字段定义（用于创建飞书多维表格）
        
        Args:
            sheet_name: Sheet 名称
            
        Returns:
            [{"name": "账号名称", "type": "text", "description": "..."}, ...]
        """
        if sheet_name not in self._wb.sheetnames:
            raise ValueError(f"Sheet '{sheet_name}' 不存在")

        ws = self._wb[sheet_name]
        fields = []

        # 第一行是字段名，第二行是字段说明
        row1 = list(next(ws.iter_rows(min_row=1, max_row=1, values_only=True)))
        row2_iter = ws.iter_rows(min_row=2, max_row=2, values_only=True)
        row2 = list(next(row2_iter, [None] * len(row1)))

        # 飞书字段类型映射
        type_map = {
            "图片": "text",       # 图片 URL 用文本
            "视频": "text",       # 视频内容用文本
            "链接": "url",        # URL 类型
            "时间": "text",       # 日期时间
            "日期": "text",
        }

        for i, name in enumerate(row1):
            if not name:
                continue
            name = str(name).strip()

            # 根据字段名推断类型
            field_type = "text"
            for keyword, ftype in type_map.items():
                if keyword in name:
                    field_type = ftype
                    break

            desc = str(row2[i]).strip() if i < len(row2) and row2[i] else ""

            fields.append({
                "name": name,
                "type": field_type,
                "description": desc,
            })

        utils.logger.info(f"[ExcelReader] 从 '{sheet_name}' 读取到 {len(fields)} 个字段定义")
        return fields

    def close(self):
        self._wb.close()

    def __enter__(self):
        return self

    def __exit__(self, *args):
        self.close()


def load_creators_from_excel(file_path: str) -> Tuple[List[Dict], List[Dict]]:
    """
    便捷函数：从 Excel 读取作者列表和字段定义
    
    Returns:
        (creators_list, result_fields)
    """
    with ExcelCreatorReader(file_path) as reader:
        creators = reader.get_creators()
        fields = reader.get_result_fields()
    return creators, fields


if __name__ == "__main__":
    # 测试
    import sys
    path = sys.argv[1] if len(sys.argv) > 1 else "redbookaccontidandresult.xlsx"
    creators, fields = load_creators_from_excel(path)
    print(f"\n作者列表 ({len(creators)}):")
    for c in creators:
        print(f"  {c['name']} | {c['id']} | {c['url'][:60]}...")
    print(f"\n结果字段 ({len(fields)}):")
    for f in fields:
        print(f"  {f['name']} ({f['type']}) - {f['description']}")
